#!/usr/bin/env python3
"""Přečte oba Excelý poolů → data/lineup-pools-import.json (pro TS import do DB).

Čte:
  data/czech-national-players.xlsx  (sheety A-tym, U20, U18, Zeny)
  data/extraliga-fantasy-players.xlsx (klubové sheety MLB…KLA; fallback Hraci)

Použití:
  python scripts/export-lineup-pools-json.py
"""

from __future__ import annotations

import json
import sys
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))

from lineup_pools_common import (  # noqa: E402
    ELH_XLSX,
    NATIONAL_XLSX,
    REPRE_POOLS,
    SHEET_TO_ELH_CLUB,
    elh_pool_key,
)
from ms2026_stable_id import stable_pool_player_id  # noqa: E402

OUT = ROOT / "data" / "lineup-pools-import.json"


def header_map(ws) -> dict[str, int]:
    row = next(ws.iter_rows(min_row=1, max_row=1, values_only=True))
    return {str(v).strip(): i for i, v in enumerate(row) if v}


def cell(row: tuple[Any, ...], headers: dict[str, int], *keys: str) -> str:
    for k in keys:
        if k in headers:
            v = row[headers[k]]
            if v is None:
                return ""
            return str(v).strip()
    return ""


def parse_jersey(raw: str) -> int | None:
    if not raw:
        return None
    try:
        n = int(float(raw))
        return n if n > 0 else None
    except ValueError:
        return None


def rows_from_sheet(ws, pool_key: str, default_club: str | None = None) -> list[dict[str, Any]]:
    headers = header_map(ws)
    if not headers:
        return []
    out: list[dict[str, Any]] = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or all(v is None or str(v).strip() == "" for v in row):
            continue
        name = cell(row, headers, "jmeno", "name")
        if not name:
            continue
        club = cell(row, headers, "klub", "club") or (default_club or "")
        position = cell(row, headers, "pozice_skupina", "position").upper()
        if position not in {"G", "D", "F"}:
            continue
        role = cell(row, headers, "pozice_detail", "role", "fantasy_pozice")
        if position == "G":
            role = "G"
        elif position == "D" and role in {"C", "LW", "RW"}:
            role = ""
        jersey = parse_jersey(cell(row, headers, "dres", "jerseyNumber", "jersey"))
        liga = cell(row, headers, "liga", "league")
        pid = cell(row, headers, "player_id", "id")
        role_key = role or position
        if not pid:
            pid = stable_pool_player_id(pool_key, name, club, role_key)
        row_out: dict[str, Any] = {
            "id": pid,
            "name": name,
            "position": position,
            "role": role or None,
            "club": club,
            "jerseyNumber": jersey,
            "poolKey": pool_key,
        }
        # Preserve Excel liga as-is (e.g. Tipsport Extraliga) — do not remap here.
        if liga and liga not in {"–", "-", "—"}:
            row_out["league"] = liga
        out.append(row_out)
    return out


def main() -> int:
    players: list[dict[str, Any]] = []
    # Never import aggregate sheets into a single pool (Všichni / Hraci without split).
    SKIP_SHEETS = {"Všichni", "Vsichni", "Legenda", "Prehled_klubu"}

    nat_path = ROOT / NATIONAL_XLSX
    if nat_path.exists():
        wb = load_workbook(nat_path, read_only=True, data_only=True)
        for name in wb.sheetnames:
            if name in SKIP_SHEETS:
                print(f"  skip national sheet {name!r} (aggregate/legend - not a pool)")
        for pool_key, sheet, _label in REPRE_POOLS:
            if sheet not in wb.sheetnames:
                print(f"WARN: missing sheet {sheet} in {nat_path.name}")
                continue
            if sheet in SKIP_SHEETS:
                continue
            # A-tym → repre_a only; never remap ELH clubs into repre_a
            chunk = rows_from_sheet(wb[sheet], pool_key)
            if pool_key == "repre_a" and len(chunk) > 220:
                raise SystemExit(
                    f"A-tym sheet produced {len(chunk)} players (expected ~171). "
                    "Refusing export — check you are not reading Všichni."
                )
            print(f"  {sheet} -> {pool_key}: {len(chunk)}")
            players.extend(chunk)
        wb.close()
    else:
        print(f"WARN: missing {nat_path}")

    elh_path = ROOT / ELH_XLSX
    if elh_path.exists():
        wb = load_workbook(elh_path, read_only=True, data_only=True)
        # Prefer short codes (MLB…) — SHEET_TO_ELH_CLUB also maps full names
        short_sheets = [
            s
            for s in wb.sheetnames
            if s in {"MLB", "LIB", "PCE", "KVA", "BRN", "LIT", "CEB", "TRI", "OLM", "PLZ", "SPA", "VIT", "MHK", "KLA"}
        ]
        if short_sheets:
            for sheet in short_sheets:
                club = SHEET_TO_ELH_CLUB[sheet]
                pool_key = elh_pool_key(club)
                chunk = rows_from_sheet(wb[sheet], pool_key, default_club=club)
                # Hard guarantee: club sheets never land in repre_a
                for row in chunk:
                    if row["poolKey"] != pool_key or not str(row["poolKey"]).startswith("elh:"):
                        raise SystemExit(f"ELH row mis-tagged poolKey={row['poolKey']!r} (sheet {sheet})")
                print(f"  ELH {sheet} -> {pool_key}: {len(chunk)}")
                players.extend(chunk)
        elif "Hraci" in wb.sheetnames:
            # fallback: split Hraci by klub — still never repre_a
            all_rows = rows_from_sheet(wb["Hraci"], "elh:__tmp__")
            by_club: dict[str, list[dict]] = {}
            for r in all_rows:
                club = r["club"]
                r["poolKey"] = elh_pool_key(club)
                r["id"] = stable_pool_player_id(r["poolKey"], r["name"], club, (r.get("role") or r["position"]))
                by_club.setdefault(club, []).append(r)
            for club, chunk in sorted(by_club.items()):
                print(f"  ELH Hraci -> {elh_pool_key(club)}: {len(chunk)}")
                players.extend(chunk)
        else:
            print("WARN: no club sheets and no Hraci in Extraliha Excel")
        wb.close()
    else:
        print(f"WARN: missing {elh_path}")

    # Dedup by id (last wins) — keep poolKey from winning row
    by_id: dict[str, dict] = {}
    for p in players:
        if not p.get("poolKey"):
            raise SystemExit(f"Player {p.get('id')} missing poolKey")
        by_id[p["id"]] = p
    deduped = list(by_id.values())

    repre_a = sum(1 for p in deduped if p["poolKey"] == "repre_a")
    elh_n = sum(1 for p in deduped if str(p["poolKey"]).startswith("elh:"))
    if repre_a == len(deduped) and elh_n == 0 and len(deduped) > 220:
        raise SystemExit(
            f"Export looks wrong: all {len(deduped)} players are repre_a. "
            "A-tym and ELH must stay separate pools."
        )

    OUT.parent.mkdir(parents=True, exist_ok=True)
    OUT.write_text(json.dumps({"players": deduped}, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(f"Wrote {OUT.relative_to(ROOT)} ({len(deduped)} players; repre_a={repre_a}, elh={elh_n})")
    return 0



if __name__ == "__main__":
    raise SystemExit(main())

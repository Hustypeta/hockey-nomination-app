#!/usr/bin/env python3
"""Export czech-extraliga-players.json → Excel pro Triple Arena / tiering."""

from __future__ import annotations

import hashlib
import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
INPUT = ROOT / "czech-extraliga-players.json"
STATS = ROOT / "data" / "extraliga-player-stats.json"
OUTPUT = ROOT / "data" / "extraliga-fantasy-players.xlsx"
DEFAULT_NATIONALITY_CODE = "CZE"
DEFAULT_NATIONALITY = "Česko"

STAT_FIELDS = [
    "gp_2025_26",
    "g_2025_26",
    "a_2025_26",
    "body_2025_26",
    "gp_2024_25",
    "g_2024_25",
    "a_2024_25",
    "body_2024_25",
    "gp_2023_24",
    "g_2023_24",
    "a_2023_24",
    "body_2023_24",
    "gp_2022_23",
    "g_2022_23",
    "a_2022_23",
    "body_2022_23",
    "gp_l10",
    "g_l10",
    "a_l10",
    "body_l10",
    "forma_body",
    "kvalita_skore",
    "tier_pismeno",
    "tier_slot",
    "tier_label",
    "v_poolu",
    "poznamka",
    "real_line",
    "formace",
    "bonus_tag",
    "bonus_ppg",
    "display_name",
]


def stable_player_id(name: str, club: str, role_key: str) -> str:
    payload = f"{name.strip()}|{club.strip()}|{role_key.strip()}"
    digest = hashlib.sha256(payload.encode("utf-8")).hexdigest()
    return f"elh_{digest[:24]}"


def fantasy_position(position: str, role: str | None) -> str:
    pos = (position or "").strip().upper()
    r = (role or "").strip().upper()
    if pos == "G":
        return "G"
    if pos == "D":
        return "D"
    if r in {"C", "LW", "RW"}:
        return r
    return ""


HEADERS = [
    ("player_id", "Stabilní ID (pro import do DB)"),
    ("jmeno", "Jméno hráče"),
    ("klub", "Klub Extraligy"),
    ("narodnost_kod", "ISO kód (CZE, SVK, …)"),
    ("narodnost", "Národnost"),
    ("pozice_skupina", "F / D / G"),
    ("pozice_detail", "C, LW, RW, RB, LB, G"),
    ("fantasy_pozice", "C / LW / RW / D / G (pro draft trojic)"),
    ("liga", "Vždy Extraliga"),
    ("formace", "Reálná lajna (1 / 2 / 3 / PP) — line stacking"),
    ("bonus_tag", "Pedigree tag (NHL) — malý bonus ve skóre"),
    ("bonus_ppg", "Volitelný ruční bonus k PPG (přepíše tag)"),
    ("gp_2025_26", "Zápasy letošní sezona"),
    ("g_2025_26", "Góly"),
    ("a_2025_26", "Asistence"),
    ("body_2025_26", "Body (G+A)"),
    ("gp_2024_25", "Zápasy 2024/25"),
    ("g_2024_25", "Góly 2024/25"),
    ("a_2024_25", "Asistence 2024/25"),
    ("body_2024_25", "Body 2024/25"),
    ("gp_2023_24", "Zápasy 2023/24"),
    ("g_2023_24", "Góly 2023/24"),
    ("a_2023_24", "Asistence 2023/24"),
    ("body_2023_24", "Body 2023/24"),
    ("gp_2022_23", "Zápasy 2022/23"),
    ("g_2022_23", "Góly 2022/23"),
    ("a_2022_23", "Asistence 2022/23"),
    ("body_2022_23", "Body 2022/23"),
    ("gp_l10", "Zápasy — forma (posledních 10)"),
    ("g_l10", "Góly L10"),
    ("a_l10", "Asistence L10"),
    ("body_l10", "Body L10"),
    ("forma_body", "Body od minulého poolu (snapshot)"),
    ("kvalita_skore", "Vypočtené skóre pro tiering — doplní skript"),
    ("tier_pismeno", "A / B / C"),
    ("tier_slot", "1=elite/solid, 2=střed, 3=risk/sleeper"),
    ("tier_label", "např. A1 Elite"),
    ("v_poolu", "ano/ne — je v aktuálním poolu 54 hráčů"),
    ("poznamka", "Poznámka"),
]

LEGEND_ROWS = [
    ("Zdroj", "czech-extraliga-players.json v projektu Lineup"),
    ("Počet hráčů", "viz list Hraci"),
    ("fantasy_pozice", "U útočníků podle role; u obránců D; u gólmanů G"),
    ("tier_pismeno + tier_slot", "Triple Arena: A1–A3, B1–B3, C1–C3"),
    ("narodnost", "Výchozí Česko (CZE); u cizinců přepsat v extraliga-player-stats.json"),
    ("formace", "Číslo lajny v Pardubicích — více hodnot oddělených čárkou (např. 1, 2)"),
    ("bonus_tag / bonus_ppg", "NHL = +0.08 PPG k tier skóre (laditelné v extraliga-tier-config.json)"),
    ("Statistiky", "Doplňuje se do data/extraliga-player-stats.json (chat / screenshoty)"),
    ("gp_l10", "Součet posledních 10 zápasů — ideálně stejná liga (ELH)"),
]


def load_stats_overlay() -> dict[str, dict]:
    if not STATS.exists():
        return {}
    raw = json.loads(STATS.read_text(encoding="utf-8"))
    return {k: v for k, v in raw.items() if isinstance(v, dict)}


def main() -> None:
    players = json.loads(INPUT.read_text(encoding="utf-8"))
    stats_by_id = load_stats_overlay()
    players.sort(key=lambda p: (p.get("club", ""), p.get("position", ""), p.get("name", "")))

    wb = Workbook()
    ws = wb.active
    ws.title = "Hraci"

    header_fill = PatternFill("solid", fgColor="1F4E79")
    header_font = Font(color="FFFFFF", bold=True)
    for col, (key, _label) in enumerate(HEADERS, start=1):
        cell = ws.cell(row=1, column=col, value=key)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for row_idx, p in enumerate(players, start=2):
        name = p.get("name", "")
        club = p.get("club", "")
        position = p.get("position", "")
        role = p.get("role") or ""
        role_key = role.strip() or position
        player_id = stable_player_id(name, club, role_key)
        overlay = stats_by_id.get(player_id, {})
        display_name = overlay.get("display_name") or name
        nat_code = overlay.get("narodnost_kod") or DEFAULT_NATIONALITY_CODE
        nat_label = overlay.get("narodnost") or DEFAULT_NATIONALITY
        values = [
            player_id,
            display_name,
            club,
            nat_code,
            nat_label,
            position,
            role,
            fantasy_position(position, role),
            "Extraliga",
            overlay.get("formace") or overlay.get("real_line", ""),
            overlay.get("bonus_tag", ""),
            overlay.get("bonus_ppg", ""),
            overlay.get("gp_2025_26", ""),
            overlay.get("g_2025_26", ""),
            overlay.get("a_2025_26", ""),
            overlay.get("body_2025_26", ""),
            overlay.get("gp_2024_25", ""),
            overlay.get("g_2024_25", ""),
            overlay.get("a_2024_25", ""),
            overlay.get("body_2024_25", ""),
            overlay.get("gp_2023_24", ""),
            overlay.get("g_2023_24", ""),
            overlay.get("a_2023_24", ""),
            overlay.get("body_2023_24", ""),
            overlay.get("gp_2022_23", ""),
            overlay.get("g_2022_23", ""),
            overlay.get("a_2022_23", ""),
            overlay.get("body_2022_23", ""),
            overlay.get("gp_l10", ""),
            overlay.get("g_l10", ""),
            overlay.get("a_l10", ""),
            overlay.get("body_l10", ""),
            overlay.get("forma_body", ""),
            overlay.get("kvalita_skore", ""),
            overlay.get("tier_pismeno", ""),
            overlay.get("tier_slot", ""),
            overlay.get("tier_label", ""),
            overlay.get("v_poolu", ""),
            overlay.get("poznamka", ""),
        ]
        for col, value in enumerate(values, start=1):
            ws.cell(row=row_idx, column=col, value=value)

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(HEADERS))}{len(players) + 1}"

    widths = {
        "A": 28,
        "B": 24,
        "C": 28,
        "D": 10,
        "E": 12,
        "F": 12,
        "G": 14,
        "H": 14,
        "I": 12,
        "J": 12,
        "K": 10,
        "L": 8,
        "M": 8,
        "N": 10,
        "O": 10,
        "P": 8,
        "Q": 8,
        "R": 12,
        "S": 10,
        "T": 8,
        "U": 8,
        "V": 12,
        "W": 10,
        "X": 8,
        "Y": 8,
        "Z": 12,
        "AA": 10,
        "AB": 8,
        "AC": 8,
        "AD": 10,
        "AE": 14,
        "AF": 16,
        "AG": 10,
        "AH": 10,
        "AI": 14,
        "AJ": 10,
        "AK": 24,
    }
    for col, width in widths.items():
        ws.column_dimensions[col].width = width

    legend = wb.create_sheet("Legenda")
    legend["A1"] = "Pole"
    legend["B1"] = "Popis"
    legend["A1"].font = Font(bold=True)
    legend["B1"].font = Font(bold=True)
    for i, (key, desc) in enumerate(HEADERS, start=2):
        legend.cell(row=i, column=1, value=key)
        legend.cell(row=i, column=2, value=desc)
    start = len(HEADERS) + 3
    legend.cell(row=start, column=1, value="—").font = Font(bold=True)
    for offset, (key, desc) in enumerate(LEGEND_ROWS, start=1):
        legend.cell(row=start + offset, column=1, value=key)
        legend.cell(row=start + offset, column=2, value=desc)
    legend.column_dimensions["A"].width = 22
    legend.column_dimensions["B"].width = 60

    clubs = wb.create_sheet("Prehled_klubu")
    clubs.append(["klub", "G", "D", "F", "celkem"])
    by_club: dict[str, dict[str, int]] = {}
    for p in players:
        club = p.get("club", "?")
        pos = p.get("position", "?")
        bucket = by_club.setdefault(club, {"G": 0, "D": 0, "F": 0})
        if pos in bucket:
            bucket[pos] += 1
    for club in sorted(by_club):
        b = by_club[club]
        total = b["G"] + b["D"] + b["F"]
        clubs.append([club, b["G"], b["D"], b["F"], total])
    clubs.freeze_panes = "A2"
    clubs.column_dimensions["A"].width = 32

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUTPUT)
    print(f"Exported {len(players)} players -> {OUTPUT}")


if __name__ == "__main__":
    main()
